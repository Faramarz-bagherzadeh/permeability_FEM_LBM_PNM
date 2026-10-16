"""Script workflows process selected inputs without custom rejection checks."""
from pathlib import Path
import weakref
import json

import numpy as np
import pytest

import run_gpu


def test_gpu_script_empty_selection_creates_no_outputs(tmp_path):
    from openpyxl import Workbook

    metadata_file = tmp_path / 'B51.xlsx'
    workbook = Workbook()
    workbook.active.append(['filename', 'porosity_open'])
    workbook.save(metadata_file)
    workbook.close()
    output = tmp_path / 'results'
    run_gpu.main([
        '--input_dir', str(tmp_path), '--input_pattern', '*.npy',
        '--output_dir', str(output), '--metadata_file', str(metadata_file),
    ])
    assert not output.exists()


def test_gpu_script_processes_tiffs_sequentially_without_input_copies(tmp_path, monkeypatch):
    import tifffile
    from porewise import load_image
    from porewise.directional import save_directional_results

    source = tmp_path / 'tif'
    source.mkdir()
    names = ['a.tif', 'b.tiff', 'c.TIFF']
    for name in names:
        image = np.ones((3, 4, 5), dtype=np.uint8)
        image[:, 1:3] = 0
        tifffile.imwrite(source / name, image, photometric='minisblack')
    (source / 'notes.txt').write_text('Ignore non-image files')
    (source / 'directory.tif').mkdir()
    output = tmp_path / 'results'
    from openpyxl import Workbook
    metadata_file = tmp_path / 'B51.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['filename', 'porosity_open'])
    for name in names:
        sheet.append([name, 0.5])
    workbook.save(metadata_file)
    workbook.close()
    monkeypatch.setattr(run_gpu, 'SAVE_EXCEL', True)
    references = []

    def load(path):
        if references:
            assert references[-1]() is None
            previous_stem = Path(names[len(references) - 1]).stem
            assert (output / f'{previous_stem}.vti').is_file()
            assert (output / f'{previous_stem}.json').is_file()
            assert (output / 'permeability_results.xlsx').is_file()
        image = load_image(path)
        references.append(weakref.ref(image))
        return image

    def save(folder, solid, **settings):
        settings.update(backend='numpy', n_steps_max=1, verbose=False)
        return save_directional_results(folder, solid, **settings)

    pytest.importorskip('openpyxl')
    monkeypatch.setattr(run_gpu, 'load_image', load)
    monkeypatch.setattr(run_gpu, 'save_directional_results', save)
    run_gpu.main([
        '--input_dir', str(source), '--input_pattern', '*',
        '--output_dir', str(output), '--metadata_file', str(metadata_file),
    ])
    assert len(references) == 3 and all(reference() is None for reference in references)
    expected_names = {'permeability_results.xlsx'}
    for name in names:
        stem = Path(name).stem
        expected_names.update({f'{stem}.json', f'{stem}.vti'})
        assert set(json.loads((output / f'{stem}.json').read_text())) == {
            'k_x_m2', 'k_y_m2', 'k_z_m2'}
    assert {path.name for path in output.iterdir()} == expected_names
    assert not list(tmp_path.rglob('*.npy'))


def test_gpu_script_filters_metadata_after_selecting_file_batch(tmp_path, monkeypatch):
    from openpyxl import Workbook

    source = tmp_path / 'samples'
    source.mkdir()
    for name in ('a.tif', 'b.tif', 'c.tif', 'd.tif'):
        (source / name).touch()
    metadata_file = tmp_path / 'B51.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['filename', 'porosity_open'])
    sheet.append(['a.tif', 0.5])
    sheet.append(['b.tif', 0.0201])
    sheet.append(['c.tif', 0.02])
    sheet.append(['d.tif', 0.0])
    workbook.save(metadata_file)
    workbook.close()
    loaded = []

    monkeypatch.setattr(run_gpu, 'load_image', lambda path: loaded.append(path.name) or np.zeros((1, 1, 1)))
    monkeypatch.setattr(run_gpu, 'save_directional_results', lambda *args, **kwargs: {
        'k_x_m2': 1.0, 'k_y_m2': 1.0, 'k_z_m2': 1.0})
    monkeypatch.setattr(run_gpu, 'SAVE_EXCEL', False)

    run_gpu.main([
        '--input_dir', str(source), '--input_pattern', '*',
        '--output_dir', str(tmp_path / 'results'),
        '--metadata_file', str(metadata_file), '--porosity_threshold', '0.02',
        '--file_offset', '1', '--file_limit', '2',
    ])

    assert loaded == ['b.tif']


def test_2d_example_reads_tiff_without_input_copy(tmp_path):
    import tifffile
    from examples.run_2d import main

    image = np.ones((6, 8), dtype=np.uint8)
    image[2:4] = 0
    path = tmp_path / 'slice.tif'
    tifffile.imwrite(path, image)
    output = tmp_path / 'results'
    assert main([str(path), '--backend', 'numpy', '--dx', '1e-6', '--steps', '1',
                 '--solid-value', '1', '--pore-value', '0', '--output', str(output)]) == 0
    assert {p.name for p in output.iterdir()} == {
        'permeability.json', 'velocity_x.npz', 'velocity_y.npz'}
    assert not list(tmp_path.rglob('*.npy'))


def test_retry_script_reads_tiffs_and_skips_completed_exports(tmp_path):
    import os
    import shutil
    import subprocess
    import sys
    import tifffile

    repository = Path(__file__).resolve().parents[1]
    script = tmp_path / 'rerun_failed_bags.sh'
    shutil.copyfile(repository / script.name, script)
    source = tmp_path / 'porewise/data/tif'
    source.mkdir(parents=True)
    for name in ('B51_bag_a.tif', 'B51_bag_b.tiff'):
        tifffile.imwrite(source / name, np.ones((3, 4, 5), dtype=np.uint8),
                         photometric='minisblack')
    # Run the actual CLI/exporter on CPU without launching a production GPU run.
    binaries = tmp_path / 'bin'
    binaries.mkdir()
    python = binaries / 'python'
    python.write_text(f'#!{sys.executable}\n' + '''import sys
if sys.argv[1:3] == ['-m', 'porewise']:
    from porewise.__main__ import main
    args = sys.argv[3:]
    args[args.index('--backend') + 1] = 'numpy'
    raise SystemExit(main(args))
if sys.argv[1] == '-':
    sys.argv = sys.argv[1:]
    exec(sys.stdin.read())
else:
    raise AssertionError(sys.argv)
''')
    python.chmod(0o755)
    env = dict(os.environ, PATH=str(binaries) + os.pathsep + os.environ['PATH'],
               PYTHONPATH=str(repository), STEPS='1')
    first = subprocess.run(['bash', str(script)], env=env, capture_output=True,
                           text=True, timeout=30)
    assert first.returncode == 0, first.stdout + first.stderr
    assert 'Finished: 2 completed, 0 failed.' in first.stdout
    second = subprocess.run(['bash', str(script)], env=env, capture_output=True,
                            text=True, timeout=30)
    assert second.returncode == 0, second.stdout + second.stderr
    assert second.stdout.count('SKIP completed:') == 2
    assert 'Starting' not in second.stdout
    assert not list(tmp_path.rglob('*.npy'))
