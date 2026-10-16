"""Process 3D image batches on the cluster using paths supplied by SLURM."""
import argparse
from pathlib import Path
import sys

from porewise import decode_mask, load_image, save_permeability_excel
from porewise.directional import save_directional_results


# Replace these example values with the values for your images.
SOLID_VALUE = 1
PORE_VALUE = 0
VOXEL_SIZE_M = 120e-6  # Voxel size in meters; e.g., 120e-6 for 120 microns.
STEPS = 20000  # Updates per direction; each image runs x, y, and z from rest.
TAU = 0.6
PRECISION = "float64"
MIRROR = False
SAVE_EXCEL = True  # Requires the optional Excel dependency.


def _open_porous_filenames(metadata_file, porosity_threshold):
    from openpyxl import load_workbook

    workbook = load_workbook(metadata_file, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        required = {"filename", "porosity_open"}
        missing = required.difference(headers)
        if missing:
            raise ValueError(
                f"Metadata file {metadata_file} is missing required columns: "
                f"{', '.join(sorted(missing))}"
            )
        filename_column = headers.index("filename")
        porosity_column = headers.index("porosity_open")
        filenames = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            filename = row[filename_column]
            porosity = row[porosity_column]
            if filename is None:
                continue
            try:
                is_open_porous = float(porosity) > porosity_threshold
            except (TypeError, ValueError):
                is_open_porous = False
            if is_open_porous:
                filenames.add(str(filename).strip())
        return filenames
    finally:
        workbook.close()


def main(argv=()):
    parser = argparse.ArgumentParser(description="Run PoreWise permeability on image volumes")
    parser.add_argument("--input_dir", type=Path, required=True)
    parser.add_argument("--input_pattern", required=True)
    parser.add_argument("--output_dir", type=Path, required=True)
    parser.add_argument("--metadata_file", type=Path, required=True)
    parser.add_argument("--porosity_threshold", type=float, default=0.02)
    parser.add_argument("--file_offset", type=int, default=0)
    parser.add_argument("--file_limit", type=int)
    args = parser.parse_args(argv)
    if args.file_offset < 0 or (args.file_limit is not None and args.file_limit < 0):
        parser.error("file_offset and file_limit must be non-negative")

    files = sorted(path for path in args.input_dir.glob(args.input_pattern)
                   if path.is_file() and path.suffix.lower() in ('.tif', '.tiff', '.npy'))
    stop = None if args.file_limit is None else args.file_offset + args.file_limit
    files = files[args.file_offset:stop]
    selected_names = _open_porous_filenames(
        args.metadata_file, args.porosity_threshold
    )
    files = [path for path in files if path.name in selected_names]

    rows = []
    for index, image_path in enumerate(files, start=1):
        print(f"[{index}/{len(files)}] {image_path.name}", flush=True)
        image = load_image(image_path)
        solid = decode_mask(image, solid_value=SOLID_VALUE, pore_value=PORE_VALUE)

        folder = args.output_dir / image_path.stem
        result = save_directional_results(
            folder,
            solid,
            voxel_size=VOXEL_SIZE_M,
            microct=image,
            backend="cuda-sparse",
            collision="trt",
            tau=TAU,
            precision=PRECISION,
            n_steps_max=STEPS,
            mirror=MIRROR,
            verbose=True,
            heartbeat=1000,
        )
        for axis in "xyz":
            print(f"k_{axis} = {result[f'k_{axis}_m2']:.6e} m²")
        print(f"Saved: {folder}", flush=True)

        if SAVE_EXCEL:
            rows.append((image_path, result))
            save_permeability_excel(args.output_dir / "permeability_results.xlsx", rows)
        del image, solid, result


if __name__ == "__main__":
    main(sys.argv[1:])
